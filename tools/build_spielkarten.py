from pathlib import Path
import json
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_ROW_HEIGHT_RULE,WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'materialien/offener-anfang';OUT.mkdir(exist_ok=True)
# Two vertically stacked cards per A4 sheet. Front/back retain top/bottom positions.
cards=[
('Stammgeflüster',1,'Zeige oder sprich. Zu zweit: Lies vor, dein Partner rät.',
 ['Die Borke schützt den Baum außen. Zeige auf das passende Wort: BORKЕ oder MARK.','Wer bin ich? A: Ich bilde neue Zellen. B: Ich transportiere Wasser nach oben. Wähle: Kambium oder Splintholz.','Spiele ohne Worte: schützen oder Wasser transportieren. Nenne danach den passenden Stammbereich. Allein: Erfinde und prüfe zwei Gesten.'],
 ['Borke. Sie liegt außen. Das Mark liegt in der Mitte.','A: Kambium. B: Splintholz.','Beispiele: Arme als Schutzschild = Borke. Hände von unten nach oben bewegen = Splintholz. Andere verständliche Gesten passen auch.']),
('Ringrennen',1,'Tippe auf eine Antwort. Zu zweit: Antworte abwechselnd.',
 ['Ein breiter Ring zeigt mehr Dickenwachstum. Welcher Ring ist breiter: 1 mm oder 4 mm?','Drei Jahresringe sind 2 mm, 4 mm und 1 mm breit. In welchem der drei Jahre war der Dickenzuwachs am größten? In welchem am kleinsten?','Erfinde eine mögliche Geschichte zum schmalen Ring. Beginne: „Vielleicht …“ Warum darfst du nicht „Ganz sicher …“ sagen?'],
 ['4 mm. Vier Millimeter sind breiter als ein Millimeter.','Am größten im zweiten Jahr mit 4 mm; am kleinsten im dritten mit 1 mm.','Zum Beispiel: „Vielleicht gab es wenig Wasser.“ Auch Licht, Temperatur und Konkurrenz wirken mit. Die Ringbreite allein beweist die Ursache nicht.']),
('Sägewerk auf Reisen',2,'Sprich die Reihenfolge. Zu zweit: Jeder ergänzt einen Schritt.',
 ['Zuerst wird der Baum gefällt. Danach wird er transportiert. Was kommt zuerst: FÄLLEN oder TRANSPORT?','Ordne diese drei Karten im Kopf: A Transport zum Sägewerk, B Baum fällen, C Stamm zu Brettern sägen. Nenne die Buchstaben.','Du bist ein Stamm. Erzähle deinen Weg vom Wald zum Brett in drei Sätzen. Baue Fällen, Transport und Einschnitt ein.'],
 ['Fällen kommt vor dem Transport.','B → A → C. Erst fällen, dann transportieren, dann sägen.','Zum Beispiel: „Ich werde gefällt. Ein Fahrzeug bringt mich zum Sägewerk. Beim Einschnitt werde ich zu Brettern gesägt.“ Andere sachlich richtige Erzählungen passen.']),
('Brett oder Pech',2,'Wähle einen Plan. Zu zweit: Spiele Kundschaft und Sägewerk.',
 ['Du brauchst ein mindestens 10 cm breites Brett. Was passt: 8 cm oder 12 cm?','Von 100 Flächenteilen werden 70 zu nutzbaren Brettern. Wie viele bleiben als Verschnitt?','Du brauchst mindestens 12 cm breite Bretter, ohne Zusammenleimen. Plan A: 75 % Ausbeute, 9 cm breit. Plan B: 65 % Ausbeute, 14 cm breit. Verkaufe der Kundschaft den passenden Plan.'],
 ['12 cm. Das ist mindestens 10 cm. 8 cm sind zu schmal.','30 Flächenteile. 100 minus 70 = 30.','Plan B liefert die nötige Breite. Plan A hat zwar mehr Ausbeute, aber seine Bretter sind zu schmal. Mehr Ausbeute allein entscheidet nicht.']),
('Laborlüge',3,'Entlarve die falsche Aussage. Zu zweit: Begründe und tausche die Rollen.',
 ['Eine Waage zeigt 50 g. Was hast du beobachtet: „Die Probe wiegt 50 g“ oder „Die Probe ist blau“?','Bei gleichem Druck hat P eine tiefe Druckspur, Q kaum eine Spur. Jemand sagt: „P ist härter.“ Stimmt das?','Behauptung: „Das schwerste Holz ist immer das härteste.“ Widerlege sie. Welchen zusätzlichen Test brauchst du?'],
 ['„Die Probe wiegt 50 g.“ Die Waage zeigt die Masse, keine Farbe.','Nein. Q widersteht dem Eindrücken stärker und ist in dieser Prüfung härter.','Masse allein belegt keine Härte. Prüfe den Widerstand gegen Eindrücken bei möglichst gleichen Bedingungen, zum Beispiel mit derselben Münze.']),
('Fair oder faul',3,'Daumen hoch = fair, Daumen runter = unfair. Begründe ab Stufe 2.',
 ['Du willst Holzarten durch Wiegen vergleichen. Fairer ist: gleich große Proben oder eine kleine und eine riesige Probe?','Beim Wassertest bekommt Probe A einen Tropfen, Probe B zehn Tropfen. Nach derselben Zeit wird verglichen. Fair oder unfair? Verbessere den Test.','Beide Proben bekommen einen Tropfen. Eine ist lackiert, die andere unbehandelt. Kannst du daraus sicher sagen, welche Holzart Wasser schneller aufnimmt?'],
 ['Gleich große Proben. Achte auch auf ähnliche Feuchte.','Unfair. Nimm dieselbe Tropfenzahl und möglichst gleich große Tropfen. Halte auch Wartezeit und weitere Bedingungen gleich.','Nein. Die Lackschicht beeinflusst den Versuch. Vergleiche ähnlich vorbereitete Oberflächen in derselben Faserrichtung; weitere Bedingungen bleiben gleich.']),
('Holz in Bewegung',4,'Zeige mit deinen Händen: größer oder kleiner. Zu zweit: Einer liest, einer spielt.',
 ['Quellen heißt: Holz wird durch Wasseraufnahme größer. Zeige Quellen: Hände auseinander oder zusammen?','Ein Brett ist erst 50 mm und nach dem Trocknen 49 mm breit. Zeige die Bewegung. Nenne den Fachbegriff.','Eine Schublade klemmt im feuchten Raum. Spiele die Schublade und erkläre eine mögliche Ursache. Nenne eine Vorsorge beim Bau.'],
 ['Hände auseinander. Das steht für größer werden.','Hände etwas zusammen. Das Brett ist geschwunden; seine Breite ist um 1 mm kleiner.','Das Holz kann Feuchte aufnehmen und quellen. Dadurch fehlt Platz. Zum Beispiel genügend Bewegungsspielraum einplanen oder passend getrocknetes Holz verwenden.']),
('Fehlerdetektive',4,'Löse die Rätsel. Zu zweit: Lies ein Rätsel vor und lass raten.',
 ['Ein Riss ist ein Spalt im Holz. Wähle: Ein Spalt heißt RISS oder AST?','A: Hier saß ein Zweig. B: Das Brett ist krumm geworden. Ordne zu: Ast oder Verwerfen.','Ein Brett hat einen kleinen, festen Ast. Jemand ruft: „Wegwerfen!“ Spiele die Gegenrede. Nenne eine mögliche Verwendung und eine Grenze.'],
 ['Riss. Ein Ast ist eine Stelle, an der ein Zweig saß.','A: Ast. B: Verwerfen.','Zum Beispiel: Ein fester Ast kann in einer dekorativen Fläche bleiben. An einer stark belasteten Verbindung muss man Lage und Zustand prüfen. Ein Ast macht Holz nicht automatisch unbrauchbar.']),
('Kundenpantomime',5,'Rate das Produkt. Zu zweit: Einer spielt, einer rät. Allein: Denke dir eine Geste aus.',
 ['Du schneidest Gemüse darauf. Was ist gesucht: SCHNEIDEBRETT oder BILDERRAHMEN?','Spiele eine Gartenbank oder ein Bücherregal. Nenne danach zwei Anforderungen an das Holz oder das Bauteil.','Dein Partner wählt: Gartenbank oder Bilderrahmen. Erkläre, warum „möglichst hart“ als einzige Anforderung nicht reicht. Allein: Vergleiche beide Produkte.'],
 ['Schneidebrett. Zum Beispiel Schneiden mit einer Handbewegung zeigen.','Gartenbank: Belastung aushalten, für draußen geeignet. Bücherregal: Bücher tragen, ausreichend steife und passend dicke Böden. Andere passende Anforderungen zählen.','Die Gartenbank wird nass und belastet. Ein Bilderrahmen soll zum Beispiel leicht sein und gut aussehen. Verwendungsort, Belastung und Gestaltung zählen zusätzlich.']),
('Wer bin ich aus Holz',5,'Nutze die Hinweise auf der Karte. Zu zweit: Einer liest vor, einer entscheidet.',
 ['Fichte ist eher weich, Buche hart. Gesucht ist das härtere Holz: FICHTE oder BUCHE?','Fichte ist leicht und preiswert, bekommt aber leicht Druckstellen. Buche ist hart und schwerer. Wähle für eine leichte, günstige Stiftebox. Begründe.','Wähle mit denselben Hinweisen für eine stark benutzte Tischfläche. Nenne einen Vorteil und einen Nachteil deiner Wahl.'],
 ['Buche ist im Vergleich das härtere Holz.','Fichte passt zu leicht und preiswert. Die möglichen Druckstellen bleiben ein Nachteil. Eine andere Wahl braucht eine passende Begründung.','Buche passt wegen der Härte besser gegen Druckstellen; das höhere Gewicht ist ein Nachteil. Fichte wäre leichter und günstiger, aber empfindlicher gegen Druckstellen.']),
('Wahrheitswächter',6,'Entscheide: wahr oder falsch. Zu zweit: Korrigiere eine falsche Aussage für deinen Partner.',
 ['Holz stammt von Bäumen. Wahr oder falsch?','A: Das Kambium ermöglicht Dickenwachstum. B: Beim Sägen entsteht nie Verschnitt. Welche Aussage ist falsch? Verbessere sie.','Erfinde zwei Aussagen über Holz: eine wahre und eine falsche. Lass raten. Allein: Schreibe beide auf einen Schmierzettel und verbessere die falsche.'],
 ['Wahr. Holz wird von Bäumen gebildet; auch Sträucher bilden Holz.','B ist falsch. Beim Sägen entsteht Verschnitt, etwa durch die Schnittbreite. A ist richtig.','Beispiel: „Holz kann quellen“ ist wahr. „Jeder Ast macht ein Brett unbrauchbar“ ist falsch. Entscheidend sind Zustand, Lage und Verwendung. Andere richtige Paare sind möglich.']),
('Knack den Holzcode',6,'Finde den Code im Kopf. Zu zweit: Einer löst, einer prüft erst danach auf der Rückseite.',
 ['Was kann Holz zum Quellen bringen? WASSER = 1. SAND = 2. Nenne die richtige Ziffer.','Nutze die Zahlen: Transport = 3, Fällen = 2, Einschnitt = 1. Ordne den Weg vom Wald zum Brett. Dein Code hat drei Ziffern.','Drei Schlösser: Härter bei gleicher Prüfung? Tiefe Spur = 4, kaum Spur = 7. Quellen? Größer = 2, kleiner = 5. Ringbreite beweist Trockenheit? Ja = 8, nein = 6. Nenne den Code und begründe jede Ziffer.'],
 ['1. Holz kann durch Wasseraufnahme in die Zellwände quellen.','231. Fällen → Transport → Einschnitt.','726. Kaum eine Spur spricht für höhere Härte. Quellen bedeutet größer werden. Ein schmaler Ring kann mehrere Ursachen haben; Trockenheit ist nicht bewiesen.'])]
# Fix a visually identical Cyrillic character in a single source token.
cards=[(t,w,m,[x.replace('BORKЕ','BORKE') for x in a],b) for t,w,m,a,b in cards]
(OUT/'spielkarten-daten.json').write_text(json.dumps([dict(id=f'S{i+1:02}',title=t,week=w,mode=m,tasks=a,solutions=b) for i,(t,w,m,a,b) in enumerate(cards)],ensure_ascii=False,indent=2))
D=Document();sec=D.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297);sec.top_margin=sec.bottom_margin=Mm(12);sec.left_margin=sec.right_margin=Mm(14);sec.header_distance=sec.footer_distance=Mm(5)
for name,size in [('Normal',12),('Title',18),('Heading 2',11)]:
 s=D.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0);s.paragraph_format.space_before=Pt(0);s.paragraph_format.space_after=Pt(5);s.paragraph_format.line_spacing=1.08
 for b in list(s.element.findall('.//'+qn('w:pBdr'))):b.getparent().remove(b)
sec.header.paragraphs[0].text='WP Technik 7   Offener Anfang   Vorderseite versuchen, dann umdrehen'
sec.header.paragraphs[0].style='Normal';sec.header.paragraphs[0].runs[0].font.size=Pt(8)
sec.footer.paragraphs[0].text='A4 · 100 % · beidseitig · lange Kante · nach dem Druck an der Trennlinie schneiden'
sec.footer.paragraphs[0].runs[0].font.size=Pt(8)
for pair in range(6):
 for reverse in [False,True]:
  if pair or reverse:D.add_page_break()
  table=D.add_table(rows=2,cols=1);table.autofit=False;table.columns[0].width=Mm(182)
  pr=table._tbl.tblPr;b=OxmlElement('w:tblBorders')
  for edge in ['top','left','bottom','right','insideH','insideV']:
   x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
  pr.append(b)
  for j in range(2):
   idx=pair*2+j;t,w,mode,tasks,sol=cards[idx];row=table.rows[j];row.height=Mm(124);row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST;row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
   c=row.cells[0];c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
   tcpr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:w'),'160');x.set(qn('w:type'),'dxa');mar.append(x)
   tcpr.append(mar)
   def p(txt='',style=None):return c.add_paragraph(txt,style)
   z=c.paragraphs[0];z.text=f'S{idx+1:02}  ·  WOCHE {w}  ·  '+('RÜCKSEITE LÖSUNGEN' if reverse else 'VORDERSEITE AUFGABEN');z.runs[0].font.size=Pt(9)
   p(t,'Title')
   if not reverse:
    p('Allein oder zu zweit · 3–5 Minuten · Wähle eine Stufe.')
    p(mode)
   else:p('Vergleiche nur die Stufe, die du ausprobiert hast. Sinngemäße Antworten zählen.')
   for k,txt in enumerate(sol if reverse else tasks):
    z=p();r=z.add_run(['1  Superleicht','2  Weiterdenken','3  Knobeln'][k]+'  ');r.bold=True;z.add_run(txt)
   z=p('Dein nächster Schritt: Erst selbst lösen. Dann umdrehen.' if not reverse else 'Dein Selbstcheck: Sicher? Probiere eine höhere Stufe. Noch unsicher? Schau nach, drehe zurück und versuche es erneut.');z.runs[0].font.size=Pt(9)
D.save(OUT/'spielkarten-offener-anfang.docx')
print('12 Karten, 24 Kartenseiten auf 12 A4-Seiten erstellt.')
